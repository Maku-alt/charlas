from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "smoke-asset.png"
OUT = ROOT / "python-pptx-smoke.pptx"

C = {
    "bg": RGBColor(0xF4, 0xF0, 0xE8),
    "ink": RGBColor(0x0F, 0x17, 0x2A),
    "muted": RGBColor(0x47, 0x55, 0x69),
    "blue": RGBColor(0x10, 0x2A, 0x43),
    "orange": RGBColor(0xC7, 0x5C, 0x2A),
    "white": RGBColor(0xFF, 0xFD, 0xF8),
    "sand": RGBColor(0xEA, 0xDB, 0xC8),
}


def ensure_asset():
    if ASSET.exists():
        return
    img = Image.new("RGB", (1200, 900), "#102A43")
    draw = ImageDraw.Draw(img)
    draw.rectangle((80, 80, 1120, 820), outline="#EADBC8", width=12)
    draw.rectangle((160, 180, 1040, 300), fill="#C75C2A")
    draw.rectangle((160, 380, 760, 470), fill="#FFFDF8")
    draw.rectangle((160, 530, 920, 620), fill="#2D6A4F")
    draw.rectangle((160, 680, 620, 740), fill="#E07A3F")
    img.save(ASSET)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def text_box(slide, text, x, y, w, h, size, color, font="Calibri", bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_header(slide, kicker, title):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = C["bg"]
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.32), Inches(1.25), Inches(0.08))
    fill(bar, C["orange"])
    text_box(slide, kicker.upper(), 0.55, 0.58, 2.8, 0.25, 9, C["orange"], "Consolas", True)
    text_box(slide, title, 0.55, 0.9, 8.9, 0.72, 29, C["ink"], "Georgia", True)


def add_footer(slide, text):
    text_box(slide, text, 0.55, 7.08, 6, 0.18, 7.5, C["muted"], "Consolas")


def build():
    ensure_asset()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = C["blue"]
    text_box(slide, "PRUEBA LOCAL", 0.65, 0.65, 2.2, 0.22, 10, RGBColor(0xE0, 0x7A, 0x3F), "Consolas", True)
    text_box(slide, "¿python-pptx abre limpio en PowerPoint nativo?", 0.65, 1.15, 8, 1.4, 34, C["white"], "Georgia", True)
    text_box(slide, "Tildes: conversación, análisis, revisión, pequeños, documentación.", 0.7, 3.02, 5.7, 0.55, 18, C["white"])
    slide.shapes.add_picture(str(ASSET), Inches(8.35), Inches(1.0), width=Inches(3.6), height=Inches(4.8))
    add_footer(slide, "python-pptx-smoke / slide 1")

    slide = prs.slides.add_slide(blank)
    add_header(slide, "comparación", "El renderer debe resolver estructura sin pedir razonamiento al LLM")
    panels = [
        ("Antes", "Agente interpreta specs, diseña, genera, revisa y corrige."),
        ("Después", "Un JSON alimenta layouts fijos y el script produce el PPTX."),
        ("Riesgo", "Si PowerPoint rechaza el archivo, el renderer no sirve para este repo."),
    ]
    for idx, (head, body) in enumerate(panels):
        x = 0.75 + idx * 4.1
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.05), Inches(3.45), Inches(2.45))
        panel.fill.solid()
        panel.fill.fore_color.rgb = C["white"]
        panel.line.color.rgb = C["sand"]
        text_box(slide, head, x + 0.28, 2.32, 2.85, 0.32, 19, C["blue"], "Georgia", True)
        text_box(slide, body, x + 0.28, 2.92, 2.78, 0.95, 14, C["ink"])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.25), Inches(11.2), Inches(0.03))
    fill(line, C["orange"])
    text_box(slide, "Criterio de aceptación: PowerPoint COM debe abrir y exportar PNG sin reparar el archivo.", 1.0, 5.48, 10.4, 0.45, 17, C["ink"], "Calibri", True)
    add_footer(slide, "python-pptx-smoke / slide 2")

    slide = prs.slides.add_slide(blank)
    add_header(slide, "cierre", "Si esta prueba falla, no insistimos con este renderer para decks nuevos")
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(4.55), Inches(13.333), Inches(2.95))
    fill(band, C["blue"])
    text_box(slide, "Decisión técnica basada en apertura nativa, no en esperanza.", 0.78, 5.18, 8.2, 0.82, 27, C["white"], "Georgia", True)
    text_box(slide, "El objetivo es bajar tokens: contenido con LLM, producción con código.", 0.82, 6.15, 7.7, 0.35, 16, C["sand"])
    slide.shapes.add_picture(str(ASSET), Inches(9.35), Inches(4.78), width=Inches(2.65), height=Inches(2.05))
    add_footer(slide, "python-pptx-smoke / slide 3")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
