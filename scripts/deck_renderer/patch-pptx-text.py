import argparse
from pathlib import Path

from pptx import Presentation


def replace_in_paragraphs(shape, old: str, new: str) -> int:
    if not hasattr(shape, "text_frame") or shape.text_frame is None:
        return 0
    count = 0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
    return count


def main() -> int:
    parser = argparse.ArgumentParser(description="Patch exact text in a PPTX without rebuilding the deck.")
    parser.add_argument("--pptx", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--slide", required=True, type=int, help="1-based slide number")
    parser.add_argument("--replace", required=True, dest="old")
    parser.add_argument("--with", required=True, dest="new")
    args = parser.parse_args()

    prs = Presentation(str(args.pptx))
    if args.slide < 1 or args.slide > len(prs.slides):
        raise SystemExit(f"slide {args.slide} out of range; deck has {len(prs.slides)} slides")

    slide = prs.slides[args.slide - 1]
    replacements = 0
    for shape in slide.shapes:
        replacements += replace_in_paragraphs(shape, args.old, args.new)

    if replacements == 0:
        raise SystemExit(f"no occurrences found on slide {args.slide}: {args.old!r}")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    print(f"patched={replacements} out={args.out.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
