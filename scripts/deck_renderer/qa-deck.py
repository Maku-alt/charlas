import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

MOJIBAKE_MARKERS = ["Ã", "Â", "�"]


def read_spec(path: Path | None) -> dict:
    if not path:
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def collect_text(pptx_path: Path) -> tuple[int, str]:
    prs = Presentation(str(pptx_path))
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                chunks.append(shape.text)
    return len(prs.slides), "\n".join(chunks)


def check_assets(spec: dict, spec_path: Path | None) -> list[str]:
    if not spec or not spec_path:
        return []
    base = spec_path.parent
    missing: list[str] = []
    for idx, slide in enumerate(spec.get("slides", []), start=1):
        image = slide.get("image")
        if not image:
            continue
        candidates = [base / image, base / "assets" / image]
        if not any(p.exists() for p in candidates):
            missing.append(f"slide {idx}: {image}")
    return missing


def check_renders(render_dir: Path | None, contact_sheet: Path | None) -> dict:
    if not render_dir:
        return {
            "render_dir": "",
            "render_png": 0,
            "blank_png": [],
            "renders_ok": True,
            "contact_sheet": "",
        }
    files = sorted([p for p in render_dir.glob("*") if p.is_file() and p.suffix.lower() == ".png"])
    blank: list[str] = []
    thumbs = []
    for p in files:
        image = Image.open(p).convert("RGB")
        if all(lo == hi for lo, hi in image.getextrema()):
            blank.append(str(p))
        if contact_sheet:
            thumb = image.copy()
            thumb.thumbnail((480, 270))
            canvas = Image.new("RGB", (500, 310), "white")
            canvas.paste(thumb, (10, 10))
            draw = ImageDraw.Draw(canvas)
            draw.text((12, 285), p.stem, fill=(15, 23, 42))
            thumbs.append(canvas)
    if contact_sheet and thumbs:
        cols = 3
        rows = (len(thumbs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * 500, rows * 310), (244, 240, 232))
        for i, thumb in enumerate(thumbs):
            sheet.paste(thumb, ((i % cols) * 500, (i // cols) * 310))
        contact_sheet.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(contact_sheet)
    return {
        "render_dir": str(render_dir.resolve()),
        "render_png": len(files),
        "blank_png": blank,
        "renders_ok": len(files) > 0 and len(blank) == 0,
        "contact_sheet": str(contact_sheet.resolve()) if contact_sheet else "",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Mechanical QA for rendered charlas PPTX decks.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--spec", type=Path, default=None)
    parser.add_argument("--expected-slides", type=int, default=None)
    parser.add_argument("--report", type=Path, default=None)
    parser.add_argument("--render-dir", type=Path, default=None)
    parser.add_argument("--contact-sheet", type=Path, default=None)
    args = parser.parse_args()

    spec = read_spec(args.spec)
    slide_count, text = collect_text(args.pptx)
    expected = args.expected_slides or len(spec.get("slides", [])) or None
    mojibake = [m for m in MOJIBAKE_MARKERS if m in text]
    missing_assets = check_assets(spec, args.spec)
    empty_text = len(text.strip()) == 0
    render_checks = check_renders(args.render_dir, args.contact_sheet)

    checks = {
        "pptx": str(args.pptx.resolve()),
        "spec": str(args.spec.resolve()) if args.spec else "",
        "slides": slide_count,
        "expected_slides": expected,
        "slide_count_ok": expected is None or slide_count == expected,
        "text_chars": len(text),
        "empty_text": empty_text,
        "mojibake_markers": mojibake,
        "mojibake_ok": len(mojibake) == 0,
        "missing_assets": missing_assets,
        "assets_ok": len(missing_assets) == 0,
        **render_checks,
    }
    checks["status"] = "ok" if all([
        checks["slide_count_ok"],
        not checks["empty_text"],
        checks["mojibake_ok"],
        checks["assets_ok"],
        checks["renders_ok"],
    ]) else "failed"

    output = json.dumps(checks, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)
    return 0 if checks["status"] == "ok" else 1


if __name__ == "__main__":
    raise SystemExit(main())
